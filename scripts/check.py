#!/usr/bin/env python3
"""Quality gate for the report. Prints RESULT: PASS or RESULT: FAIL.
  python3 scripts/check.py            offline checks
  python3 scripts/check.py --online   also confirms every source link answers today
Checks:
  1  every fact quote is found word for word in its saved snapshot, and every snapshot matches its SHA-256 in evidence/manifest.json
  2  every citation points to a known fact, every source is cited, every paragraph with a status label cites at least one fact
  3  no wording that suggests a partnership, talks, plans or an agreement between Redbelly and Quant
  4  no price talk, investment advice or hype words, no en dash, em dash or emoji
  5  no personal name or handle from the local, git-ignored .credit-names file anywhere in the deliverables
  6  each X post fits in 280 characters
  7  the PDF, the Word file and the web page contain the current text of every paragraph
"""
import hashlib, json, re, subprocess, sys, urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
errors, notes = [], []
R = json.loads((ROOT / "content/report.json").read_text())
FACTS = json.loads((ROOT / "content/facts.json").read_text()); F = {f["id"]: f for f in FACTS}
SOURCES = {s["id"]: s for s in json.loads((ROOT / "content/sources.json").read_text())}
CITE = re.compile(r"\{([a-z0-9 -]+)\}")

def norm(s):
    for a, b in (("’", "'"), ("‘", "'"), ("“", '"'), ("”", '"'), (" ", " "), (" ", " ")):
        s = s.replace(a, b)
    return re.sub(r"\s+", " ", s).strip()

# 1 quotes and snapshots
man = {m["file"]: m for m in json.loads((ROOT / "evidence/manifest.json").read_text())}
for fpath, m in man.items():
    p = ROOT / fpath
    if not p.exists():
        errors.append(f"evidence: {fpath} listed in manifest but missing")
    elif hashlib.sha256(p.read_bytes()).hexdigest() != m["sha256"]:
        errors.append(f"evidence: {fpath} changed after capture (SHA-256 mismatch)")
for s in SOURCES.values():
    rel = str((ROOT / "evidence/snapshots" / s["file"]).resolve().relative_to(ROOT))
    if rel not in man:
        errors.append(f"source {s['id']}: snapshot {rel} not in manifest")
for f in FACTS:
    s = SOURCES.get(f["source"])
    if not s:
        errors.append(f"fact {f['id']}: unknown source {f['source']}"); continue
    snap = (ROOT / "evidence/snapshots" / s["file"]).read_text(encoding="utf-8", errors="ignore")
    if norm(f["quote"]) not in norm(snap):
        errors.append(f"fact {f['id']}: quote not found in {s['file']}")

# 2 citations and labels
prose, labelled = [], []   # prose: (where, text) written by us; labelled: blocks that need a citation
prose.append(("status note", R["status_note"]))
for k in R["key_findings"]:
    prose += [("finding", k["q"]), ("finding", k["a"] + " {" + k["cite"] + "}")]
for s in R["sections"]:
    prose.append((s["id"] + " title", s["title"]))
    if s.get("intro"):
        prose.append((s["id"] + " intro", s["intro"]))
    for i, b in enumerate(s["blocks"]):
        where = f"{s['id']} block {i + 1}"
        if b["label"] not in ("confirmed", "tested", "analysis", "unknown", "summary_table", "method"):
            errors.append(f"{where}: unknown label {b['label']}")
        parts = [b.get("heading", ""), b.get("text", "")] + b.get("list", []) + [r[0] + " " + r[1] + " {" + r[2] + "}" for r in b.get("rows", [])]
        for t in parts:
            if t:
                prose.append((where, t))
        if b["label"] in ("confirmed", "tested", "analysis", "unknown"):
            for t in ([b["text"]] if b.get("text") else []) + b.get("list", []):
                if not CITE.search(t):
                    errors.append(f"{where}: a {b['label']} statement cites no source: {t[:60]}")
summary = [("community summary", p) for p in R["community_summary"]["paragraphs"]] + [("X thread", t) for t in R["community_summary"]["thread"]]
used = set()
for where, t in prose:
    for m in CITE.finditer(t):
        for fid in m.group(1).split():
            if fid not in F:
                errors.append(f"{where}: cites unknown fact {fid}")
            used.add(fid)
for fid in F:
    if fid not in used:
        notes.append(f"fact {fid} is not cited (kept in the fact list only)")
for sid in SOURCES:
    if not any(F[f]["source"] == sid for f in used if f in F):
        errors.append(f"source {sid} is never cited")
if R["key_findings"] and not all(k["label"] in ("confirmed", "analysis", "unknown", "tested") for k in R["key_findings"]):
    errors.append("key findings: every row needs a status label")

# 3 partnership wording (our own words only; source quotes are checked in 1)
texts = [(w, CITE.sub("", t)) for w, t in prose + summary]
NEG = re.compile(r"\b(no|not|none|nothing|never|without)\b", re.I)
for w, t in texts:
    for sent in re.split(r"(?<=[.?!])\s+", t):
        if sent.rstrip().endswith("?"):
            continue
        if re.search(r"partner|talks\b|negotiat|collaborat|agreement|in discussion|working with|plans?\b|roadmap|will (integrate|connect|join|support|list)", sent, re.I) \
           and re.search(r"redbelly|quant", sent, re.I) and not NEG.search(sent):
            if not re.search(r"The Clearing House", sent):
                errors.append(f"{w}: wording could suggest a Redbelly and Quant link: {sent[:90]}")

# 4 price talk, advice, hype, dashes, emoji
BANNED = r"\b(price|prices|priced|profit|profits|returns?|yield|apy|apr|buy|sell|bullish|bearish|moon|pump|undervalued|investment advice|financial advice|guaranteed|revolutionary|game.?changer|unleash|seamless|skyrocket|massive)\b"
EMOJI = re.compile("[\U0001F300-\U0001FAFF☀-➿⬀-⯿️]")
for w, t in texts:
    m = re.search(BANNED, t, re.I)
    if m:
        errors.append(f"{w}: banned word '{m.group(0)}'")
for w, t in texts + [("fact " + f["id"], f["quote"]) for f in FACTS]:
    if "–" in t or "—" in t:
        errors.append(f"{w}: contains an en dash or em dash")
    if EMOJI.search(t):
        errors.append(f"{w}: contains an emoji or symbol")

# 5 personal credit
cn = ROOT / ".credit-names"
deliverables = [p for p in ROOT.rglob("*") if p.is_file() and not any(x in p.parts for x in (".git", "evidence", "build", "__pycache__")) and p.name != ".credit-names"]
if cn.exists():
    names = [n.strip().lower() for n in cn.read_text().splitlines() if n.strip()]
    for p in deliverables:
        if p.suffix in (".woff2",):
            continue
        if p.suffix == ".pdf":
            body = subprocess.run(["pdftotext", str(p), "-"], capture_output=True, text=True).stdout
        elif p.suffix == ".docx":
            from docx import Document
            body = "\n".join(x.text for x in Document(p).paragraphs)
        else:
            body = p.read_bytes().decode("utf-8", "ignore")
        for n in names:
            if n in body.lower():
                errors.append(f"personal credit: {p.relative_to(ROOT)} contains a listed name")
    if re.search(r"made by|created by|author:", " ".join(t for _, t in texts), re.I):
        errors.append("personal credit: a 'made by' style line")
else:
    notes.append("personal credit scan skipped, .credit-names not present")

# 6 X thread length (a link counts as 23 characters on X)
for i, t in enumerate(R["community_summary"]["thread"], 1):
    n = len(t.replace("{site}", "x" * 23))
    if n > 280:
        errors.append(f"X post {i} is {n} characters, over 280")

# 7 outputs carry the current text
def plain(t):
    return norm(CITE.sub("", t))
pdf = ROOT / "public/downloads/redbelly-overledger-feasibility-report.pdf"
docx = ROOT / "public/downloads/redbelly-overledger-feasibility-report.docx"
web = ROOT / "public/index.html"
outs = {}
if pdf.exists():
    outs["PDF"] = norm(subprocess.run(["pdftotext", "-raw", str(pdf), "-"], capture_output=True, text=True).stdout)
if docx.exists():
    from docx import Document
    d = Document(docx)
    outs["Word file"] = norm("\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells]))
if web.exists():
    import html as H
    outs["web page"] = norm(H.unescape(re.sub(r"<[^>]+>", " ", web.read_text())))
for name in ("PDF", "Word file", "web page"):
    if name not in outs:
        errors.append(f"{name} missing, run scripts/build.py")
        continue
    o = re.sub(r"\[\s*\d+(\s*,\s*\d+)*\s*\]", "", outs[name]); o = re.sub(r"\s+", " ", o)
    for where, t in prose + summary:
        probe = re.sub(r"\s+", " ", plain(t))[:70].strip()
        probe = probe.replace("{site}", "")
        if probe and re.sub(r"\s", "", probe) not in re.sub(r"\s", "", o):
            errors.append(f"{name} is out of date: '{probe[:50]}' ({where}) not found, run scripts/build.py")
            break

# online: every source link answers
if "--online" in sys.argv:
    for s in SOURCES.values():
        if s.get("local"):
            continue
        try:
            req = urllib.request.Request(s["url"], headers={"User-Agent": "Mozilla/5.0"})
            code = urllib.request.urlopen(req, timeout=30).status
        except Exception as e:
            code = getattr(e, "code", str(e)[:40])
        if code != 200:
            notes.append(f"link {s['id']}: answered {code} (check in a browser: {s['url']})")
        else:
            notes.append(f"link {s['id']}: 200")

for n in notes:
    print("note:", n)
for e in errors:
    print("FAIL:", e)
print(f"{len(FACTS)} facts, {len(SOURCES)} sources, {len(man)} evidence files checked")
print("RESULT:", "PASS" if not errors else f"FAIL ({len(errors)})")
sys.exit(1 if errors else 0)
